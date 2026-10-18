import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


def load_env(path: Path):
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def request(method: str, url: str, key: str, body=None, prefer=None):
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    try:
        with urllib.request.urlopen(
            urllib.request.Request(url, data=payload, headers=headers, method=method),
            timeout=60,
        ) as response:
            return response.read().decode("utf-8")
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase returned {error.code}: {detail}") from error


def exact_count(url: str, key: str):
    request_object = urllib.request.Request(
        f"{url}&limit=1",
        headers={
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Prefer": "count=exact",
        },
        method="GET",
    )
    with urllib.request.urlopen(request_object, timeout=60) as response:
        content_range = response.headers.get("Content-Range", "")
    return int(content_range.rsplit("/", 1)[1])


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text and not text.startswith("=DISPIMG(") else None


def normalise_part_number(value: str):
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def main():
    repo = Path(__file__).resolve().parents[1]
    load_env(repo / ".env.local")
    workbook_path = Path(sys.argv[1])
    supabase_url = os.environ["SUPABASE_URL"].rstrip("/")
    service_key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    endpoint = f"{supabase_url}/rest/v1/part_cross_reference"
    source = "Seven Brands"
    imported_at = datetime.now(timezone.utc).isoformat()

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rows = []
    brand_counts = {}
    for sheet in workbook.worksheets:
        headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        positions = {name: index for index, name in enumerate(headers) if name}
        required = {"MIC Code", "Part Number"}
        if not required.issubset(positions):
            raise RuntimeError(f"{sheet.title} is missing required columns")

        count = 0
        for values in sheet.iter_rows(min_row=2, values_only=True):
            mic_code = clean(values[positions["MIC Code"]])
            part_number = clean(values[positions["Part Number"]])
            if not mic_code or not part_number:
                continue
            rows.append({
                "brand": sheet.title,
                "model": clean(values[positions["Model"]]) if "Model" in positions else None,
                "mic_code": mic_code,
                "part_number": part_number,
                "part_number_normalized": normalise_part_number(part_number),
                "chinese_name": clean(values[positions["Chinese Name"]]) if "Chinese Name" in positions else None,
                "description": clean(values[positions["Description"]]) if "Description" in positions else None,
                "source": source,
                "source_updated_at": imported_at,
            })
            count += 1
        brand_counts[sheet.title] = count

    source_filter = urllib.parse.quote(source, safe="")
    request("DELETE", f"{endpoint}?source=eq.{source_filter}", service_key)
    for start in range(0, len(rows), 500):
        request("POST", endpoint, service_key, rows[start:start + 500], "return=minimal")

    verified_rows = exact_count(
        f"{endpoint}?select=id&source=eq.{source_filter}",
        service_key,
    )
    print(json.dumps({
        "imported": len(rows),
        "chinese_names": sum(1 for row in rows if row["chinese_name"]),
        "brands": brand_counts,
        "verified_rows": verified_rows,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
